# -*- coding: utf-8 -*-
"""
Stage 3 of the email-to-chat skill.

Render a segmented conversation as a chat-style .docx:

    python render_chat_docx.py <chat.json> <output.docx> [images-dir]

chat.json schema
----------------
{
  "title": "Project X - kickoff",
  "messages": [
    { "sender": "Irene Perini",
      "date":   "2026-06-18",          ISO, used for the date separators
      "time":   "14:59",               as printed in the message header
      "domain": "liu.se",              sender's email domain, drives the side split
      "blocks": [
        {"type": "text",  "value": "Dear all,"},
        {"type": "ref",   "value": "pubmed.ncbi.nlm.nih.gov/28213645/"},
        {"type": "image", "value": "image001.jpg"},
        {"type": "note",  "value": "[2 images - not carried over]"}
      ] } ] }

Colour and side are decided here, not by the caller: bubble colours are
assigned from a fixed palette in order of first appearance, and the side split
follows the dominant email domain. See assign_colours() / decide_sides().
"""

import json
import os
import sys
from datetime import date as _date

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Segoe UI"
BG_PAGE = "EDE4DB"
DATE_PILL = "D6D0C8"
TXT = RGBColor(0x11, 0x1B, 0x21)
META = RGBColor(0x60, 0x74, 0x7A)
SYS = RGBColor(0x7A, 0x70, 0x66)

# (bubble fill, sender-name colour) assigned in order of first appearance
PALETTE = [
    ("DCEEFB", "145C8F"),   # blue
    ("FFF1D6", "9C5A0E"),   # amber
    ("D9FDD3", "1F6B36"),   # green
    ("EEE3F7", "5B3A8C"),   # lilac
    ("FFE2E2", "A03434"),   # rose
    ("DFF5F3", "126B66"),   # teal
    ("F5EEDC", "7A6224"),   # sand
    ("E6E6F5", "3E3E8C"),   # periwinkle
]

BUBBLE_W = Inches(4.85)
IMAGE_W = Inches(2.9)


# --------------------------------------------------------------------------
# low-level docx helpers
# --------------------------------------------------------------------------

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_margins(cell, top=90, start=140, bottom=90, end=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start),
                     ("bottom", bottom), ("end", end)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    table._tbl.tblPr.append(borders)


def cant_split(table):
    """Keep a bubble on one page. Word ignores this if the row exceeds a page."""
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def page_background(doc, hexcolor):
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), hexcolor)
    doc.element.insert(0, bg)
    doc.settings.element.append(OxmlElement("w:displayBackgroundShape"))


def spacer(doc, pts):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)
    p.add_run("").font.size = Pt(2)


def text_par(cell, text, size=10.5, color=TXT, italic=False,
             indent=0.0, space_after=3):
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.08
    pf.left_indent = Inches(indent)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.italic = italic
    return p


# --------------------------------------------------------------------------
# colour / side assignment
# --------------------------------------------------------------------------

def assign_colours(messages):
    """Palette entry per sender, in order of first appearance."""
    order = []
    for m in messages:
        if m["sender"] not in order:
            order.append(m["sender"])
    return {name: PALETTE[i % len(PALETTE)] for i, name in enumerate(order)}, order


def decide_sides(messages, senders):
    """
    Two-sided layout only while it still carries meaning: the domain that sent
    the most messages goes left, everyone else right. With more than four
    participants or more than two domains a two-sided split stops being
    informative and starts misleading, so everything goes left.
    """
    domains = {}
    for m in messages:
        d = (m.get("domain") or "").lower()
        domains[d] = domains.get(d, 0) + 1
    domains.pop("", None)

    if len(senders) > 4 or len(domains) > 2 or len(domains) < 2:
        return {name: "left" for name in senders}, False

    main = max(domains, key=lambda d: domains[d])
    sides = {}
    for m in messages:
        d = (m.get("domain") or "").lower()
        sides[m["sender"]] = "left" if d == main else "right"
    for name in senders:
        sides.setdefault(name, "left")
    return sides, True


def pretty_date(iso):
    try:
        d = _date.fromisoformat(iso)
    except (ValueError, TypeError):
        return str(iso).upper()
    return "%d %s %d" % (d.day, d.strftime("%B").upper(), d.year)


# --------------------------------------------------------------------------
# blocks
# --------------------------------------------------------------------------

def date_separator(doc, label):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_borders(t)
    t.autofit = False
    w = Inches(1.9)
    t.columns[0].width = w
    c = t.cell(0, 0)
    c.width = w
    shade(c, DATE_PILL)
    cell_margins(c, 50, 90, 50, 90)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label)
    r.font.name = FONT
    r.font.size = Pt(8.5)
    r.bold = True
    r.font.color.rgb = SYS
    spacer(doc, 5)


def bubble(doc, msg, fill, name_hex, side, images_dir, missing):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = (WD_TABLE_ALIGNMENT.RIGHT if side == "right"
                   else WD_TABLE_ALIGNMENT.LEFT)
    no_borders(t)
    t.autofit = False
    t.columns[0].width = BUBBLE_W
    c = t.cell(0, 0)
    c.width = BUBBLE_W
    shade(c, fill)
    cell_margins(c)
    cant_split(t)

    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(msg["sender"])
    r.font.name = FONT
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(name_hex)

    for b in msg.get("blocks", []):
        kind, val = b.get("type", "text"), b.get("value", "")
        if kind == "text":
            text_par(c, val)
        elif kind == "note":
            text_par(c, val, size=9, color=META, italic=True)
        elif kind == "ref":
            text_par(c, val, size=9.5, indent=0.18, space_after=2)
        elif kind == "image":
            path = os.path.join(images_dir, val) if images_dir else val
            if not os.path.isfile(path):
                missing.append(val)
                text_par(c, "[image not found: %s]" % val,
                         size=9, color=META, italic=True)
                continue
            pp = c.add_paragraph()
            pp.paragraph_format.space_before = Pt(3)
            pp.paragraph_format.space_after = Pt(3)
            pp.add_run().add_picture(path, width=IMAGE_W)

    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(msg.get("time", ""))
    r.font.name = FONT
    r.font.size = Pt(7.5)
    r.font.color.rgb = META

    spacer(doc, 7)


def legend(doc, senders, colours):
    """Colour key, up to four participants per row."""
    per_row = 4
    for start in range(0, len(senders), per_row):
        chunk = senders[start:start + per_row]
        t = doc.add_table(rows=1, cols=len(chunk))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(t)
        t.autofit = False
        for i, name in enumerate(chunk):
            fill, name_hex = colours[name]
            c = t.cell(0, i)
            c.width = Inches(1.7)
            shade(c, fill)
            cell_margins(c, 40, 80, 40, 80)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(name)
            r.font.name = FONT
            r.font.size = Pt(8.5)
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(name_hex)
        spacer(doc, 4)
    spacer(doc, 10)


# --------------------------------------------------------------------------

def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__.strip())
    chat_path, out_path = sys.argv[1], sys.argv[2]
    images_dir = sys.argv[3] if len(sys.argv) > 3 else \
        os.path.join(os.path.dirname(os.path.abspath(chat_path)), "images")

    with open(chat_path, encoding="utf-8") as fh:
        chat = json.load(fh)
    messages = chat.get("messages", [])
    if not messages:
        sys.exit("chat.json contains no messages")

    colours, senders = assign_colours(messages)
    sides, two_sided = decide_sides(messages, senders)

    doc = Document()
    page_background(doc, BG_PAGE)

    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(chat.get("title", "Conversation"))
    r.font.name = FONT
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = TXT

    legend(doc, senders, colours)

    missing = []
    last_date = None
    for m in messages:
        if m.get("date") and m["date"] != last_date:
            date_separator(doc, pretty_date(m["date"]))
            last_date = m["date"]
        fill, name_hex = colours[m["sender"]]
        bubble(doc, m, fill, name_hex, sides[m["sender"]], images_dir, missing)

    doc.save(out_path)

    n_img = sum(1 for m in messages for b in m.get("blocks", [])
                if b.get("type") == "image") - len(missing)
    print("written      %s" % out_path)
    print("messages     %d" % len(messages))
    print("layout       %s" % ("two-sided" if two_sided else "single column"))
    for name in senders:
        print("  %-24s %s  %s" % (name, colours[name][0], sides[name]))
    print("images       %d embedded" % n_img)
    for mi in missing:
        print("WARNING: image not found: %s" % mi)


if __name__ == "__main__":
    main()
